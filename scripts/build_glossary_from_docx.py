#!/usr/bin/env python
"""Build glossary/brand_glossary.tbx from a client guidelines .docx Brand Lexicon table.

Usage: python scripts/build_glossary_from_docx.py <guidelines.docx> [--table N]

Finds the EN|FR lexicon table (or the one given with --table), splits multi-line cells
into aligned pairs, applies a small set of explicit normalisations for source typos /
usage notes, writes TBX-Basic, and prints an anomalies list for human confirmation.
"""
from __future__ import annotations

import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "glossary" / "brand_glossary.tbx"

# Explicit, transparent normalisations (row source FR -> (clean term, note)).
# Keeps the glossary functional (QA does substring matching) while surfacing the
# original so the client can confirm/correct their source of truth.
NORMALISE = {
    "Honoraires (toujours au pluriel)": ("Honoraires", "always plural (toujours au pluriel)"),
    "Chapiter / section": ("Chapitre", "source read 'Chapiter'; alt: 'section'"),
    "Groupe des hauts revenus (High Income Earners GroupO":
        ("Groupe des hauts revenus", "source had a stray '(High Income Earners GroupO'"),
    "Entreprise de conférenciers.ères":
        ("Entreprise de conférencier·ères", "normalised '.' to median dot '·'"),
}


def find_lexicon_table(doc: Document, forced: int | None):
    if forced is not None:
        return doc.tables[forced]
    for t in doc.tables:
        if len(t.columns) == 2 and len(t.rows) > 5:
            hdr = [c.text.strip().lower() for c in t.rows[0].cells]
            if hdr[0] in ("term", "english", "en") and "fr" in hdr[1] or hdr[1] == "french":
                return t
    raise SystemExit("Could not auto-find a 2-column lexicon table; pass --table N")


def split_pairs(en: str, fr: str) -> list[tuple[str, str]]:
    en_lines = [x.strip() for x in en.split("\n") if x.strip()]
    fr_lines = [x.strip() for x in fr.split("\n") if x.strip()]
    if len(en_lines) == len(fr_lines) and len(en_lines) > 1:
        return list(zip(en_lines, fr_lines))
    return [(" ".join(en_lines), " ".join(fr_lines))]


def main() -> int:
    args = sys.argv[1:]
    forced = None
    if "--table" in args:
        i = args.index("--table")
        forced = int(args[i + 1])
        args = args[:i] + args[i + 2:]
    if not args:
        raise SystemExit("give the guidelines .docx path")

    doc = Document(args[0])
    table = find_lexicon_table(doc, forced)

    entries: list[tuple[str, str, str]] = []  # (en, fr, note)
    anomalies: list[str] = []
    for r in table.rows[1:]:  # skip header
        en_raw = r.cells[0].text.strip()
        fr_raw = r.cells[1].text.strip()
        if not en_raw or not fr_raw:
            continue
        for en, fr in split_pairs(en_raw, fr_raw):
            note = ""
            if fr in NORMALISE:
                clean, note = NORMALISE[fr]
                anomalies.append(f"{en!r}: {fr!r} -> {clean!r} ({note})")
                fr = clean
            entries.append((en, fr, note))

    # Emit TBX-Basic.
    parts = ['<?xml version="1.0" encoding="UTF-8"?>',
             '<!-- Generated from client guidelines by build_glossary_from_docx.py -->',
             '<martif type="TBX-Basic" xml:lang="en"><martifHeader><fileDesc>',
             '<sourceDesc><p>CAPS French Language Guidelines — Brand Lexicon</p></sourceDesc>',
             '</fileDesc></martifHeader><text><body>']
    for i, (en, fr, note) in enumerate(entries, 1):
        note_xml = f'<descrip type="note">{escape(note)}</descrip>' if note else ""
        parts.append(
            f'<termEntry id="c{i}">{note_xml}'
            f'<langSet xml:lang="en"><tig><term>{escape(en)}</term></tig></langSet>'
            f'<langSet xml:lang="fr"><tig><term>{escape(fr)}</term></tig></langSet>'
            f'</termEntry>'
        )
    parts.append('</body></text></martif>')
    OUT.write_text("\n".join(parts), encoding="utf-8")

    print(f"Wrote {len(entries)} terms -> {OUT}")
    if anomalies:
        print("\nAnomalies normalised (confirm with the client):")
        for a in anomalies:
            print("  -", a)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
