"""Assemble a bilingual DOCX as sequential blocks (CAPS-compliant).

Full first-language document, a page break + visual divider, then the full
second-language document — in ONE .docx, preserving all formatting of both (headings,
tables, images) via docxcompose. Order defaults to English-first; French-first is the
CAPS rule for Quebec materials.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxcompose.composer import Composer

from .config import Config


def _insert_notice_at_top(doc: Document, text: str) -> None:
    """Add a centered, italic bilingual notice as the first paragraph."""
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
    # Move the newly-appended paragraph to the top of the body.
    doc.element.body.insert(0, p._p)


def _append_divider(doc: Document, text: str) -> None:
    """Start a new page and add a centered, bold divider heading."""
    doc.add_page_break()
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.bold = True


def combine(cfg: Config, en_docx: str | Path, fr_docx: str | Path,
            out_docx: str | Path, order: str | None = None) -> Path:
    """Produce a bilingual DOCX. order: 'en-fr' (default) or 'fr-en'."""
    b = cfg.raw.get("bilingual", {})
    order = order or b.get("default_order", "en-fr")
    if order not in ("en-fr", "fr-en"):
        raise ValueError("order must be 'en-fr' or 'fr-en'")

    if order == "en-fr":
        first, second = Path(en_docx), Path(fr_docx)
        notice = b.get("notice_en_first", "")
        divider = b.get("divider_fr", "———")
    else:
        first, second = Path(fr_docx), Path(en_docx)
        notice = b.get("notice_fr_first", "")
        divider = b.get("divider_en", "———")

    master = Document(str(first))
    if notice:
        _insert_notice_at_top(master, notice)
    _append_divider(master, divider)

    composer = Composer(master)
    composer.append(Document(str(second)))

    out = Path(out_docx)
    out.parent.mkdir(parents=True, exist_ok=True)
    composer.save(str(out))
    return out
