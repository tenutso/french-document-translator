"""Generate tests/fixtures/sample.docx — a small multi-feature manual page.

Includes headings, styled runs, a numbered list, a table, a footer, and an embedded
image containing English text, so the round-trip, QA, and image stages all have
something realistic to exercise on the pod.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw

FIX = Path(__file__).parent / "fixtures"


def _make_image(path: Path) -> None:
    img = Image.new("RGB", (480, 160), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([2, 2, 477, 157], outline="black", width=2)
    d.text((20, 40), "Figure 1: Click the Save button", fill="black")
    d.text((20, 90), "then choose Export to email.", fill="black")
    img.save(path)


def main() -> None:
    FIX.mkdir(parents=True, exist_ok=True)
    img_path = FIX / "_fig1.png"
    _make_image(img_path)

    doc = Document()
    doc.add_heading("Getting Started", level=1)
    p = doc.add_paragraph("Welcome to the training manual. This section explains how to ")
    p.add_run("log in").bold = True
    p.add_run(" and check your email for 2 important messages.")

    doc.add_heading("Step-by-step", level=2)
    for step in ["Open the application.", "Enter your credentials.",
                 "Empty the shopping cart before you continue."]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("Reference table", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Action"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "Download the file"
    table.cell(1, 1).text = "The file is saved locally."

    doc.add_heading("Screenshot", level=2)
    doc.add_picture(str(img_path))

    section = doc.sections[0]
    section.footer.paragraphs[0].text = "Confidential — Page footer"

    for style_name in ("Normal",):
        doc.styles[style_name].font.size = Pt(11)

    out = FIX / "sample.docx"
    doc.save(str(out))
    img_path.unlink(missing_ok=True)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
