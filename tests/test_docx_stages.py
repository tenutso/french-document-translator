"""Tests for XLIFF I/O, image location resolution, and report rendering.

These avoid external tools (Okapi/Java, vLLM, tesseract binary) so they run anywhere.
Okapi extract/merge and CometKiwi are exercised on the pod (see README verification).
"""
from pathlib import Path

from qc_translate.config import load_config
from qc_translate.images import extract_images
from qc_translate.models import ImageInfo, Segment
from qc_translate.report import write_image_report, write_qa_report
from qc_translate.xliff import read_sources, write_targets

CFG = load_config()
FIX = Path(__file__).parent / "fixtures"

MINI_XLIFF = """<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
 <file source-language="en" target-language="fr-ca" datatype="x-docx" original="source.docx">
  <body>
   <trans-unit id="1"><source>Click <g id="1">Save</g> now.</source></trans-unit>
   <trans-unit id="2"><source>Keep this <x id="3"/> tag.</source></trans-unit>
   <trans-unit id="3" translate="no"><source>DO_NOT_TOUCH</source></trans-unit>
  </body>
 </file>
</xliff>"""


def test_read_sources_skips_translate_no(tmp_path: Path):
    xlf = tmp_path / "m.xlf"
    xlf.write_text(MINI_XLIFF, encoding="utf-8")
    units = read_sources(xlf)
    ids = [u for u, _ in units]
    assert ids == ["1", "2"]  # unit 3 (translate=no) skipped
    assert 'g id="1"' in units[0][1]


def test_write_targets_preserves_inline_and_reads_back(tmp_path: Path):
    xlf = tmp_path / "m.xlf"
    xlf.write_text(MINI_XLIFF, encoding="utf-8")
    out = tmp_path / "out.xlf"
    write_targets(xlf, {
        "1": 'Cliquez sur <g id="1">Enregistrer</g> maintenant.',
        "2": 'Conservez cette <x id="3"/> balise.',
    }, out)
    text = out.read_text(encoding="utf-8")
    assert "<target" in text and "Enregistrer" in text
    assert 'id="1"' in text and 'id="3"' in text  # inline codes survived


def test_write_targets_malformed_falls_back_to_text(tmp_path: Path):
    xlf = tmp_path / "m.xlf"
    xlf.write_text(MINI_XLIFF, encoding="utf-8")
    out = tmp_path / "out.xlf"
    # Broken markup should not crash the writer.
    write_targets(xlf, {"1": "Texte <g id=1 sans fermeture"}, out)
    assert out.exists()


def test_image_location_resolution():
    # Uses the generated fixture; OCR is unavailable locally but location must resolve.
    imgs = extract_images(CFG, FIX / "sample.docx", translate_fn=None)
    assert len(imgs) == 1
    im = imgs[0]
    assert im.media_path == "word/media/image1.png"
    assert "Screenshot" in im.location  # nearest preceding heading
    assert im.thumbnail_data_uri.startswith("data:image/png;base64,")


def _tmp_cfg(tmp_path: Path):
    cfg = load_config()
    cfg.raw["tm"]["db"] = str(tmp_path / "tm.sqlite")
    cfg.raw["tm"]["tmx_export"] = str(tmp_path / "tm.tmx")
    return cfg


def test_import_review_xliff_updates_tm(tmp_path: Path):
    from qc_translate.review import import_review
    from qc_translate.tm import TranslationMemory
    job = tmp_path / "job"; job.mkdir()
    (job / "source.docx.xlf").write_text(MINI_XLIFF, encoding="utf-8")
    reviewed = tmp_path / "reviewed.xlf"
    reviewed.write_text(MINI_XLIFF.replace(
        "<trans-unit id=\"1\"><source>Click <g id=\"1\">Save</g> now.</source></trans-unit>",
        "<trans-unit id=\"1\"><source>Click <g id=\"1\">Save</g> now.</source>"
        "<target>Cliquez sur <g id=\"1\">Enregistrer</g> maintenant.</target></trans-unit>"),
        encoding="utf-8")
    cfg = _tmp_cfg(tmp_path)
    updated, total = import_review(cfg, reviewed, job)
    assert updated == 1
    tm = TranslationMemory(cfg.tm["db"])
    assert "Enregistrer" in (tm.exact("Click Save now.") or "")
    tm.close()


def test_package_builds_zip(tmp_path: Path):
    import zipfile
    from docx import Document
    from qc_translate.review import package
    job = tmp_path / "job"; job.mkdir()
    Document().save(str(job / "sample.fr-CA.draft.docx"))
    (job / "qa_report.html").write_text("<html>qa</html>")
    cfg = _tmp_cfg(tmp_path)
    zip_path, files = package(cfg, job)
    names = zipfile.ZipFile(zip_path).namelist()
    assert "REVIEW_INSTRUCTIONS.md" in names
    assert "sample.fr-CA.draft.docx" in names and "qa_report.html" in names
    # CAT-tool interchange files: TM (TMX) + glossary (TBX) travel with the package so a
    # reviewer can import them into Smartcat / OmegaT.
    assert "qc_translate.tmx" in names
    assert "brand_glossary.tbx" in names


def test_finalize_docx_sets_language_and_fields(tmp_path: Path):
    import zipfile
    from docx import Document
    from qc_translate.okapi import finalize_docx

    p = tmp_path / "d.docx"
    d = Document(); d.add_paragraph("Bonjour le monde."); d.save(str(p))
    finalize_docx(p, target_lang="fr-CA")

    z = zipfile.ZipFile(p)
    settings = z.read("word/settings.xml").decode()
    styles = z.read("word/styles.xml").decode()
    document = z.read("word/document.xml").decode()
    assert "updateFields" in settings and 'w:val="true"' in settings
    assert "fr-CA" in settings          # themeFontLang
    assert "fr-CA" in styles            # docDefaults lang -> French proofing
    assert 'w:val="fr-CA"' in document  # explicit per-run lang (survives concatenation)
    # still a valid, readable docx
    assert Document(str(p)).paragraphs[0].text == "Bonjour le monde."


def test_reports_render(tmp_path: Path):
    segs = [
        Segment("1", "Send the email.", "Envoyez le courriel.", qe_score=0.91),
        Segment("2", "Open <g id='1'>file</g>.", "Ouvrir.", qa_flags=["tag_mismatch"],
                qe_score=0.42),
    ]
    qa = write_qa_report(tmp_path / "qa.html", "job", segs, 0.75)
    html = qa.read_text()
    assert "tag_mismatch" in html and "0.42" in html

    imgs = [ImageInfo("word/media/image1.png", "rId1", "under heading X",
                      ocr_text="Save", has_text=True, requires_recreation=True,
                      suggested_fr="Enregistrer")]
    ir = write_image_report(tmp_path / "img.html", "job", imgs)
    assert "Enregistrer" in ir.read_text()
